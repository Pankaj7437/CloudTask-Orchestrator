import docx
import re
import os

doc_path = "CloudTask_AoA_Project_Report.docx"
print(f"=== COMPREHENSIVE AUDIT OF {doc_path} ===")

doc = docx.Document(doc_path)

# 1. Total statistics
total_paragraphs = len(doc.paragraphs)
total_tables = len(doc.tables)

full_text = []
for p in doc.paragraphs:
    full_text.append(p.text)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                full_text.append(p.text)

complete_corpus = " ".join(full_text)
words = complete_corpus.split()
total_words = len(words)

print(f"1. Paragraphs: {total_paragraphs}")
print(f"2. Tables: {total_tables}")
print(f"3. Total Words: {total_words}")

# 2. Images audit
rels = doc.part.rels
image_rels = [r for r in rels.values() if "image" in r.target_ref]
print(f"4. Total Embedded Images: {len(image_rels)}")
for r in image_rels:
    print(f"   -> rId: {r.rId}, Target: {r.target_ref}")

# Check images referenced in document body
referenced_rids = set()
for p in doc.paragraphs:
    xml = p._p.xml
    matches = re.findall(r'r:embed="([^"]+)"', xml)
    referenced_rids.update(matches)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                matches = re.findall(r'r:embed="([^"]+)"', p._p.xml)
                referenced_rids.update(matches)

print(f"5. Referenced Image IDs in Document: {len(referenced_rids)}")
for rid in referenced_rids:
    if rid in rels:
        print(f"   -> {rid} : {rels[rid].target_ref}")

# 3. Keyword Checks
checks = {
    "Asharam": False,
    "Ratnaram": True,
    "internship": False,
    "Industrial Training": False,
    "Learn and Build": False,
    "LnB": False,
    "PASTE UI SCREENSHOT": False,
    "placeholder": False,
    "2026-27": True,
    "5th Semester": True,
    "Pankaj Sharma": True,
    "24EGJCS147": True,
    "Global Institute of Technology": True,
    "Dr. Pradeep Jha": True,
    "Dr. I. C. Sharma": True
}

print("\n=== KEYWORD & METADATA VERIFICATION ===")
all_pass = True
for term, expected in checks.items():
    found_count = len(re.findall(re.escape(term), complete_corpus, re.IGNORECASE))
    status = (found_count > 0) == expected
    mark = "PASS" if status else "FAIL"
    if not status:
        all_pass = False
    print(f"[{mark}] Term '{term}': found {found_count} times (Expected: {'Present' if expected else 'Absent'})")

# 4. Check for Figure Captions
figures_found = re.findall(r'Figure \d+\.\d+:', complete_corpus)
print(f"\n6. Figures Captioned in Document: {len(figures_found)}")
for fig in figures_found:
    print(f"   -> {fig}")

# 5. Check Page Breaks
page_breaks = 0
for p in doc.paragraphs:
    if 'w:br' in p._p.xml and 'w:type="page"' in p._p.xml:
        page_breaks += 1
print(f"\n7. Total Hard Page Breaks: {page_breaks}")

print("\n" + ("="*50))
if all_pass and len(referenced_rids) >= 13:
    print("ALL VERIFICATION CHECKS PASSED PERFECTLY!")
else:
    print("WARNING: SOME CHECKS NEED ATTENTION.")
print("="*50)
